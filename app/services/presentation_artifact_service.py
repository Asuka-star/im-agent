import json
import logging
import re
from html import escape
from pathlib import Path
from typing import Any

from app.services.artifact_skills import ArtifactVerifier, SlidesSkill
from app.services.artifact_skills.style_tokens import ARTIFACT_STYLE, css_color, pptx_color
from app.utils.values import coerce_positive_int


logger = logging.getLogger(__name__)


class PresentationArtifactService:
    GENERIC_TITLES = {"presentation", "slides", "slide deck", "ppt", "汇报大纲", "演示稿", "ppt演示稿"}

    def __init__(self, *, root_dir: Path | None = None) -> None:
        self.root_dir = root_dir or Path("data") / "artifacts" / "slides"
        self.slides_skill = SlidesSkill()
        self.verifier = ArtifactVerifier()

    def persist_package(
        self,
        package: dict,
        *,
        provider: str,
        task_run_id: str | None,
        session_id: str,
    ) -> dict:
        self.root_dir.mkdir(parents=True, exist_ok=True)
        preview = self._normalize_package(package)
        preview["theme"] = self._package_title(preview)
        json_filename = self._slides_filename(preview, task_run_id=task_run_id, session_id=session_id)
        html_filename = self._html_filename(json_filename)
        pptx_filename = self._pptx_filename(json_filename)
        pdf_filename = self._pdf_filename(json_filename)
        preview["exports"] = {
            "json": f"/api/artifacts/slides/{json_filename}",
            "html": f"/api/artifacts/slides/{html_filename}",
            "pptx": f"/api/artifacts/slides/{pptx_filename}",
        }

        pptx_path = self.root_dir / pptx_filename
        self._write_pptx(pptx_path, preview)
        pptx_check = self.verifier.verify_pptx_file(pptx_path)
        if not pptx_check.ok:
            raise ValueError(f"Generated pptx failed verification: {pptx_check.warnings}")
        self._attach_pdf_export_best_effort(preview, pdf_filename)

        (self.root_dir / json_filename).write_text(
            json.dumps(preview, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (self.root_dir / html_filename).write_text(self._render_html(preview), encoding="utf-8")

        return {
            "artifact_type": "slides_package",
            "provider": provider,
            "status": "ready",
            "title": self._package_title(preview),
            "url": f"/api/artifacts/slides/{html_filename}",
            "preview": preview,
            "version": coerce_positive_int(preview.get("version")),
        }

    def _normalize_package(self, package: dict) -> dict:
        return self.slides_skill.normalize(package)

    @classmethod
    def _package_title(cls, package: dict) -> str:
        theme = str(package.get("theme") or "").strip()
        if theme and theme.lower() not in cls.GENERIC_TITLES:
            return theme
        slides = package.get("slides") if isinstance(package.get("slides"), list) else []
        for slide in slides:
            if not isinstance(slide, dict):
                continue
            title = str(slide.get("title") or "").strip()
            if title and title.lower() not in cls.GENERIC_TITLES:
                return f"{title}演示稿" if "演示" not in title and "汇报" not in title else title
        return "协作演示稿"

    def _slides_filename(self, package: dict, *, task_run_id: str | None, session_id: str) -> str:
        title = self._package_title(package)
        stem = self._artifact_stem(title, suffix=task_run_id or session_id)[:96]
        return f"{stem or 'slides'}.json"

    def _html_filename(self, json_filename: str) -> str:
        stem = json_filename.rsplit(".", 1)[0]
        return f"{stem or 'slides'}.html"

    def _pptx_filename(self, json_filename: str) -> str:
        stem = json_filename.rsplit(".", 1)[0]
        return f"{stem or 'slides'}.pptx"

    def _pdf_filename(self, json_filename: str) -> str:
        stem = json_filename.rsplit(".", 1)[0]
        return f"{stem or 'slides'}.pdf"

    def _artifact_stem(self, title: str, *, suffix: str | None) -> str:
        title_stem = self._slugify_filename(title)
        suffix_stem = self._short_suffix(suffix)
        if suffix_stem and suffix_stem not in title_stem:
            return f"{title_stem}-{suffix_stem}"
        return title_stem or suffix_stem or "slides"

    @staticmethod
    def _short_suffix(value: str | None) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        if text.startswith("run_") and len(text) > 24:
            return text[:16]
        return text[:24]

    def _attach_pdf_export_best_effort(self, package: dict, pdf_filename: str) -> None:
        try:
            self._write_pdf(self.root_dir / pdf_filename, package)
        except Exception as exc:  # noqa: BLE001
            logger.warning("Failed to export presentation PDF, keeping other exports: %s", exc)
            warnings = package.setdefault("export_warnings", [])
            if isinstance(warnings, list):
                warnings.append({"export": "pdf", "error": str(exc)})
            return
        exports = package.setdefault("exports", {})
        if isinstance(exports, dict):
            exports["pdf"] = f"/api/artifacts/slides/{pdf_filename}"

    def _slugify_filename(self, value: str) -> str:
        slug = re.sub(r"[^\w.-]+", "-", value.strip(), flags=re.UNICODE).strip(".-_")
        return slug or "artifact"

    def _render_html(self, package: dict) -> str:
        theme = str(package.get("theme") or "Presentation").strip()
        audience = str(package.get("audience") or "").strip()
        slides = package.get("slides") if isinstance(package.get("slides"), list) else []
        slide_html = "\n".join(self._render_slide(slide, index) for index, slide in enumerate(slides, start=1))
        audience_html = f"<p>{escape(audience)}</p>" if audience else ""
        return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{escape(theme)}</title>
  <style>
    :root {{ color-scheme: light; font-family: {ARTIFACT_STYLE["font"]}; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: {css_color("background")}; color: {css_color("ink")}; }}
    header {{ padding: 24px 32px 18px; background: {css_color("primary_dark")}; color: white; border-bottom: 5px solid {css_color("accent")}; }}
    header h1 {{ margin: 0; font-size: 28px; line-height: 1.2; letter-spacing: 0; }}
    header p {{ margin: 8px 0 0; color: #D4E7EA; }}
    main {{ display: grid; gap: 22px; padding: 24px; max-width: 1120px; margin: 0 auto; }}
    section {{ aspect-ratio: 16 / 9; background: {css_color("surface")}; border: 1px solid {css_color("border")}; border-radius: 8px; padding: 42px; display: grid; grid-template-rows: auto auto 1fr auto; gap: 18px; box-shadow: 0 14px 34px rgba(20, 92, 100, .12); }}
    .index {{ font-size: 13px; color: {css_color("primary")}; font-weight: 800; text-transform: uppercase; }}
    h2 {{ margin: 0; color: {css_color("ink")}; font-size: clamp(28px, 4vw, 44px); line-height: 1.14; letter-spacing: 0; }}
    ul {{ margin: 0; padding: 0; list-style: none; display: grid; gap: 12px; font-size: clamp(17px, 2vw, 23px); line-height: 1.4; }}
    li {{ position: relative; padding-left: 24px; }}
    li::before {{ content: ""; position: absolute; left: 0; top: .62em; width: 8px; height: 8px; border-radius: 50%; background: {css_color("accent")}; }}
    aside {{ margin-top: 4px; background: {css_color("accent_soft")}; border-left: 4px solid {css_color("accent")}; border-radius: 6px; padding: 12px 14px; color: {css_color("ink")}; font-size: 14px; line-height: 1.45; }}
    aside strong {{ color: {css_color("primary")}; }}
  </style>
</head>
<body>
  <header>
    <h1>{escape(theme)}</h1>
    {audience_html}
  </header>
  <main>
    {slide_html}
  </main>
</body>
</html>
"""

    def _render_slide(self, slide: object, index: int) -> str:
        payload = slide if isinstance(slide, dict) else {}
        title = str(payload.get("title") or f"Slide {index}").strip()
        bullets = payload.get("bullets") if isinstance(payload.get("bullets"), list) else []
        notes = str(payload.get("speaker_notes") or "").strip()
        duration_sec = payload.get("duration_sec")
        bullet_html = "\n".join(f"<li>{escape(str(item))}</li>" for item in bullets[:5]) or "<li>No key points yet</li>"
        note_html = ""
        if notes:
            note_html = (
                f"<aside><strong>Speaker notes</strong> | {escape(str(duration_sec or 45))}s<br>"
                f"{escape(notes)}</aside>"
            )
        return f"""<section>
  <div class="index">Slide {index:02d}</div>
  <h2>{escape(title)}</h2>
  <ul>{bullet_html}</ul>
  {note_html}
</section>"""

    def _write_pptx(self, path: Path, package: dict) -> None:
        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.enum.text import PP_ALIGN
            from pptx.util import Inches, Pt
        except ImportError as exc:  # pragma: no cover - exercised only when dependency is missing.
            raise RuntimeError("python-pptx is required to export pptx files") from exc

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]
        slides = package.get("slides") if isinstance(package.get("slides"), list) else []
        if not slides:
            slides = [{"title": package.get("theme") or "Presentation", "bullets": ["No key points yet"]}]

        for index, slide_payload in enumerate(slides, start=1):
            slide = prs.slides.add_slide(blank_layout)
            payload = slide_payload if isinstance(slide_payload, dict) else {}
            title = str(payload.get("title") or f"Slide {index}").strip()
            bullets = payload.get("bullets") if isinstance(payload.get("bullets"), list) else []
            notes = str(payload.get("speaker_notes") or "").strip()

            background = slide.background.fill
            background.solid()
            background.fore_color.rgb = RGBColor.from_string(pptx_color("background"))

            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor.from_string(pptx_color("accent"))
            band.line.fill.background()

            self._add_textbox(
                slide,
                text=f"{index:02d}",
                left=0.62,
                top=0.55,
                width=0.8,
                height=0.28,
                size=11,
                color=pptx_color("primary"),
                bold=True,
                align=PP_ALIGN.LEFT,
            )
            self._add_textbox(
                slide,
                text=title,
                left=0.62,
                top=0.9,
                width=11.5,
                height=1.02,
                size=30,
                color=pptx_color("ink"),
                bold=True,
                align=PP_ALIGN.LEFT,
            )

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.62), Inches(2.05), Inches(12.05), Inches(4.2))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor.from_string(pptx_color("surface"))
            card.line.color.rgb = RGBColor.from_string(pptx_color("border"))
            card.line.width = Pt(1)

            self._add_bullets(
                slide,
                bullets=[str(item).strip() for item in bullets if str(item).strip()] or ["No key points yet"],
                left=1.0,
                top=2.42,
                width=11.15,
                height=3.32,
            )
            self._add_textbox(
                slide,
                text=str(package.get("theme") or "Presentation"),
                left=0.66,
                top=6.76,
                width=9.8,
                height=0.28,
                size=9,
                color=pptx_color("muted"),
                align=PP_ALIGN.LEFT,
            )
            self._add_textbox(
                slide,
                text=f"{index}/{len(slides)}",
                left=11.82,
                top=6.76,
                width=0.8,
                height=0.28,
                size=9,
                color=pptx_color("muted"),
                align=PP_ALIGN.RIGHT,
            )

            if notes:
                slide.notes_slide.notes_text_frame.text = notes

        prs.save(path)

    def _add_textbox(
        self,
        slide: Any,
        *,
        text: str,
        left: float,
        top: float,
        width: float,
        height: float,
        size: int,
        color: str,
        bold: bool = False,
        align: Any = None,
    ) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        if align is not None:
            paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = RGBColor.from_string(color)

    def _add_bullets(self, slide: Any, *, bullets: list[str], left: float, top: float, width: float, height: float) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        for index, bullet in enumerate(bullets[:5]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_after = Pt(9)
            run = paragraph.add_run()
            run.text = f"- {bullet}"
            run.font.size = Pt(18)
            run.font.name = "Microsoft YaHei"
            run.font.color.rgb = RGBColor.from_string(pptx_color("ink"))

    def _write_pdf(self, path: Path, package: dict) -> None:
        slides = package.get("slides") if isinstance(package.get("slides"), list) else []
        if not slides:
            slides = [{"title": package.get("theme") or "Presentation", "bullets": ["No key points yet"]}]
        theme = str(package.get("theme") or "Presentation")
        page_streams = [
            self._pdf_page_stream(theme, slide if isinstance(slide, dict) else {}, index, len(slides))
            for index, slide in enumerate(slides, start=1)
        ]
        path.write_bytes(self._build_pdf(page_streams))

    def _pdf_page_stream(self, theme: str, slide: dict, index: int, total: int) -> bytes:
        title = str(slide.get("title") or f"Slide {index}").strip()
        bullets = slide.get("bullets") if isinstance(slide.get("bullets"), list) else []
        notes = str(slide.get("speaker_notes") or "").strip()
        duration = str(slide.get("duration_sec") or 45)
        ops: list[str] = [
            "0.96 0.98 0.97 rg 0 0 842 595 re f",
            "0.08 0.35 0.40 rg 0 0 842 54 re f",
            f"BT /F1 14 Tf 42 570 Td {self._pdf_hex_text(theme)} Tj ET",
            f"BT /F1 12 Tf 742 570 Td {self._pdf_hex_text(f'{index}/{total}')} Tj ET",
            f"BT /F1 28 Tf 42 500 Td {self._pdf_hex_text(title)} Tj ET",
        ]
        y = 448
        for bullet in [str(item).strip() for item in bullets if str(item).strip()][:5]:
            for line_index, line in enumerate(self._wrap_text(bullet, 32)[:2]):
                prefix = "- " if line_index == 0 else "  "
                ops.append(f"BT /F1 17 Tf 66 {y} Td {self._pdf_hex_text(prefix + line)} Tj ET")
                y -= 27
        if notes:
            y = min(y - 20, 268)
            ops.extend(
                [
                    "0.90 0.96 0.96 rg 42 60 758 180 re f",
                    "0.09 0.41 0.48 RG 42 60 758 180 re S",
                    f"BT /F1 13 Tf 62 214 Td {self._pdf_hex_text(f'讲者备注 | {duration}s')} Tj ET",
                ]
            )
            note_y = 186
            for line in self._wrap_text(notes, 46)[:5]:
                ops.append(f"BT /F1 12 Tf 62 {note_y} Td {self._pdf_hex_text(line)} Tj ET")
                note_y -= 23
        return ("\n".join(ops) + "\n").encode("ascii")

    def _build_pdf(self, page_streams: list[bytes]) -> bytes:
        objects: list[bytes] = [
            b"<< /Type /Catalog /Pages 5 0 R >>",
            (
                b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light "
                b"/Encoding /UniGB-UCS2-H /DescendantFonts [3 0 R] >>"
            ),
            (
                b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
                b"/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 2 >> "
                b"/FontDescriptor 4 0 R /DW 1000 >>"
            ),
            (
                b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 6 "
                b"/FontBBox [0 -200 1000 900] /ItalicAngle 0 /Ascent 800 /Descent -200 "
                b"/CapHeight 700 /StemV 80 >>"
            ),
        ]
        page_object_ids: list[int] = []
        next_id = 6
        for stream in page_streams:
            page_id = next_id
            content_id = next_id + 1
            page_object_ids.append(page_id)
            objects.append(
                (
                    f"<< /Type /Page /Parent 5 0 R /MediaBox [0 0 842 595] "
                    f"/Resources << /Font << /F1 2 0 R >> >> /Contents {content_id} 0 R >>"
                ).encode("ascii")
            )
            objects.append(b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"endstream")
            next_id += 2
        kids = " ".join(f"{item} 0 R" for item in page_object_ids)
        objects.insert(4, f"<< /Type /Pages /Kids [{kids}] /Count {len(page_object_ids)} >>".encode("ascii"))

        output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = [0]
        for index, payload in enumerate(objects, start=1):
            offsets.append(len(output))
            output.extend(f"{index} 0 obj\n".encode("ascii"))
            output.extend(payload)
            output.extend(b"\nendobj\n")
        xref_offset = len(output)
        output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        output.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        output.extend(
            (
                f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
                f"startxref\n{xref_offset}\n%%EOF\n"
            ).encode("ascii")
        )
        return bytes(output)

    def _pdf_hex_text(self, value: str) -> str:
        return "<" + str(value or "").encode("utf-16-be", errors="ignore").hex().upper() + ">"

    def _wrap_text(self, value: str, width: int) -> list[str]:
        text = " ".join(str(value or "").split())
        if not text:
            return []
        lines: list[str] = []
        while text:
            lines.append(text[:width])
            text = text[width:]
        return lines
